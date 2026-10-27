"""End-to-end check on the synthetic sample set:  python -m pytest tests  (or python tests/test_samples.py)"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.dirname(__file__))

import openpyxl  # noqa: E402

import make_samples  # noqa: E402
from isoqc.cli import main  # noqa: E402


def _run():
    tmp = tempfile.mkdtemp()
    make_samples.main(tmp)
    out = os.path.join(tmp, "out")
    assert main(["--base", tmp, "--out", out, "--workers", "1", "--no-pause"]) == 0
    wb = openpyxl.load_workbook(os.path.join(out, "ISO_QC_Report.xlsx"))
    summary = {r[1]: r[2] for r in wb["요약"].iter_rows(min_row=2, values_only=True)}
    detail = {}
    for r in wb["상세"].iter_rows(min_row=2, values_only=True):
        detail.setdefault(r[0], []).append((r[4], r[8] or "", r[9] or "", r[11]))
    return summary, detail


def test_samples():
    summary, detail = _run()
    assert summary["ISO-1TC-PL-2302-30-024-01"] == "육안확인"      # all text reflected, one empty cloud
    assert summary["ISO-1TC-PL-2302-30-025-01"] == "NG"
    assert summary["ISO-1TC-PL-2302-30-026-01"] == "As-built 없음"

    ok = {(old, new): st for _, old, new, st in detail["ISO-1TC-PL-2302-30-024-01"]}
    ng = {(old, new): st for _, old, new, st in detail["ISO-1TC-PL-2302-30-025-01"]}
    expected_ng = {("SOUTH", "WEST"), ("2610", "2618"), ("5", "10"), ("N", "Y")}
    for key in [("NORTH", "SOUTH"), ("SOUTH", "WEST"), ("2610", "2618"), ("Open", "CLOSED"),
                ("5", "10"), ("N", "Y")]:
        assert ok[key] == "반영", (key, ok)
        assert ng[key] == ("미반영" if key in expected_ng else "반영"), (key, ng)
    # SHX helper annotations and the signature never become checked items
    assert all(st != "미반영" for _, _, _, st in detail["ISO-1TC-PL-2302-30-024-01"])


if __name__ == "__main__":
    test_samples()
    print("OK")
