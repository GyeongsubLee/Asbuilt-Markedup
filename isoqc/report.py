"""Write the QC report: Excel workbook + HTML pages with evidence images."""
from __future__ import annotations

import html
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config as C
from .process import verdict

FILL = {
    C.OK: "C6EFCE", "OK": "C6EFCE",
    C.NG: "FFC7CE", "NG": "FFC7CE", "오류": "FFC7CE", "As-built 없음": "FFC7CE",
    C.CHECK: "FFEB9C", "확인필요": "FFEB9C",
    C.VISUAL: "DDEBF7", "육안확인": "DDEBF7",
    C.INFO: "EDEDED", "Mark-up 없음": "EDEDED",
}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(color="FFFFFF", bold=True)


def _rel(path: str, out_dir: str) -> str:
    return os.path.relpath(path, out_dir).replace(os.sep, "/") if path else ""


def _counts(items, region=None):
    st = [i["status"] for i in items if region is None or i["region"] == region]
    return {s: st.count(s) for s in (C.OK, C.NG, C.CHECK, C.VISUAL, C.INFO)}


def _region_verdict(items, region):
    c = _counts(items, region)
    if c[C.NG]:
        return f"NG ({c[C.NG]})"
    if c[C.CHECK]:
        return f"확인필요 ({c[C.CHECK]})"
    if c[C.VISUAL]:
        return f"육안확인 ({c[C.VISUAL]})"
    if c[C.OK]:
        return f"OK ({c[C.OK]})"
    return "-"


def _sheet(ws, headers, widths):
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def _style_row(ws, row, status_col=None, status=None, wrap_cols=()):
    for cell in ws[row]:
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=cell.column in wrap_cols)
    if status_col and status in FILL:
        ws.cell(row=row, column=status_col).fill = PatternFill("solid", fgColor=FILL[status])


def write_excel(results, markup_only, asbuilt_only, dups, out_dir, meta) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    _sheet(ws, ["No", "ISO No", "종합판정", "1. Routing 반영", "2. Line Condition 반영", "BOM",
                "반영", "미반영", "확인필요", "육안확인", "참고", "Mark-up 파일", "As-built 파일",
                "정렬 품질", "비고", "상세"],
           [6, 30, 12, 16, 18, 12, 7, 7, 8, 8, 7, 34, 34, 30, 30, 10])
    for n, r in enumerate(results, 1):
        v = verdict(r)
        c = _counts(r["items"])
        ws.append([n, r["iso"], v, _region_verdict(r["items"], "Routing"),
                   _region_verdict(r["items"], "Line Condition"), _region_verdict(r["items"], "BOM"),
                   c[C.OK], c[C.NG], c[C.CHECK], c[C.VISUAL], c[C.INFO],
                   os.path.basename(r["markup"]), os.path.basename(r["asbuilt"] or ""),
                   r.get("align", ""), r.get("error", ""), "보기"])
        row = ws.max_row
        _style_row(ws, row, 3, v, wrap_cols=(15,))
        link = ws.cell(row=row, column=16)
        link.hyperlink = f"html/{r['html']}"
        link.font = Font(color="0563C1", underline="single")
    for k, p in markup_only:
        ws.append([ws.max_row, k, "As-built 없음", "", "", "", "", "", "", "", "", os.path.basename(p), "",
                   "", "AS-BUILT ISO 폴더에 같은 ISO No 파일 없음", ""])
        _style_row(ws, ws.max_row, 3, "As-built 없음")

    wd = wb.create_sheet("상세")
    _sheet(wd, ["ISO No", "Mark-up 파일", "#", "Page", "구분", "항목(Condition)", "Mark-up 유형", "주석 종류",
                "기존값", "Mark-up 값", "As-built 값", "판정", "사유", "Mark-up 이미지", "As-built 이미지"],
           [28, 30, 5, 5, 13, 22, 16, 18, 22, 24, 30, 9, 40, 12, 12])
    for r in results:
        for i in r["items"]:
            wd.append([r["iso"], os.path.basename(r["markup"]), i["no"], i["page"], i["region"], i["field"],
                       i["kind"], i["source"], i["old"], i["new"], i["asbuilt"], i["status"], i["reason"],
                       "열기" if i["img_mk"] else "", "열기" if i["img_ab"] else ""])
            row = wd.max_row
            _style_row(wd, row, 12, i["status"], wrap_cols=(9, 10, 11, 13))
            for col, key in ((14, "img_mk"), (15, "img_ab")):
                if i[key]:
                    c = wd.cell(row=row, column=col)
                    c.hyperlink = _rel(i[key], out_dir)
                    c.font = Font(color="0563C1", underline="single")
        if r.get("error"):
            wd.append([r["iso"], os.path.basename(r["markup"]), "", "", "", "", "", "", "", "", "", "오류",
                       r["error"]])
            _style_row(wd, wd.max_row, 12, "오류")

    wm = wb.create_sheet("미매칭")
    _sheet(wm, ["구분", "ISO No", "파일"], [30, 32, 70])
    for k, p in markup_only:
        wm.append(["Mark-up 만 있음 (As-built 없음)", k, p])
    for k, p in asbuilt_only:
        wm.append(["As-built 만 있음 (Mark-up 없음)", k, p])
    for k, ps in dups.items():
        for p in ps:
            wm.append(["같은 ISO No 파일 여러 개", k, p])

    wi = wb.create_sheet("정보")
    for k, v in meta.items():
        wi.append([k, v])
    wi.column_dimensions["A"].width = 22
    wi.column_dimensions["B"].width = 90
    wi.append([])
    wi.append(["판정 기준", ""])
    for k, v in (
        (C.OK, "Mark-up 값이 As-built 같은 위치(칸/라인)에 있음"),
        (C.NG, "As-built 에 Mark-up 값이 없거나 기존값이 그대로 남아 있음"),
        (C.CHECK, "값은 있으나 위치가 다르거나 기존값도 근처에 있어 판단이 애매함"),
        (C.VISUAL, "문자 없는 구름/도형(형상 변경) - 이미지로 육안 확인"),
        (C.INFO, "서명/스탬프/Title·Revision 영역 - 검토 대상 아님"),
    ):
        wi.append([k, v])
        wi.cell(row=wi.max_row, column=1).fill = PatternFill("solid", fgColor=FILL[k])

    path = os.path.join(out_dir, "ISO_QC_Report.xlsx")
    wb.save(path)
    return path


# --------------------------------------------------------------------------- HTML
CSS = """
:root{--bg:#f6f7f9;--card:#fff;--ink:#1b1f24;--mut:#5b6573;--line:#d9dde3;
--ok:#1e7b34;--okb:#e3f4e6;--ng:#b42318;--ngb:#fde8e7;--ck:#8a6100;--ckb:#fff4cf;
--vi:#1f5aa6;--vib:#e4eefb;--in:#5b6573;--inb:#eef0f3}
@media (prefers-color-scheme:dark){:root{--bg:#14171b;--card:#1d2127;--ink:#e8eaed;--mut:#9aa4b1;
--line:#2f353d;--okb:#163222;--ok:#7fd69a;--ngb:#3b1a18;--ng:#ff9b92;--ckb:#3a2f10;--ck:#f3cf6a;
--vib:#172a44;--vi:#8fb8f0;--inb:#262b32;--in:#aab3bf}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);
font:14px/1.45 "Malgun Gothic","Apple SD Gothic Neo",system-ui,sans-serif}
main{max-width:1500px;margin:0 auto;padding:20px 16px}h1{font-size:20px;margin:0 0 4px}
.sub{color:var(--mut);margin-bottom:16px}.cards{display:flex;gap:10px;flex-wrap:wrap;margin-bottom:16px}
.card{background:var(--card);border:1px solid var(--line);border-radius:8px;padding:10px 14px;min-width:110px}
.card b{display:block;font-size:22px}table{border-collapse:collapse;width:100%;background:var(--card)}
th,td{border:1px solid var(--line);padding:6px 8px;text-align:left;vertical-align:top}
th{background:var(--inb);position:sticky;top:0}.wrap{overflow-x:auto}
.s{display:inline-block;padding:1px 8px;border-radius:10px;font-weight:600;white-space:nowrap}
.s-반영,.s-OK{background:var(--okb);color:var(--ok)}.s-미반영,.s-NG,.s-오류{background:var(--ngb);color:var(--ng)}
.s-확인필요{background:var(--ckb);color:var(--ck)}.s-육안확인{background:var(--vib);color:var(--vi)}
.s-참고,.s-Mark-up{background:var(--inb);color:var(--in)}
img{max-width:420px;width:100%;border:1px solid var(--line);border-radius:4px;background:#fff}
.imgs{display:grid;grid-template-columns:1fr 1fr;gap:8px;min-width:560px}.imgs small{color:var(--mut)}
input{padding:6px 10px;border:1px solid var(--line);border-radius:6px;background:var(--card);color:var(--ink);
width:min(360px,100%);margin-bottom:10px}a{color:var(--vi)}
"""


def _badge(s):
    return f'<span class="s s-{html.escape(s.split()[0])}">{html.escape(s)}</span>'


def write_html(results, out_dir, meta) -> str:
    hdir = os.path.join(out_dir, "html")
    os.makedirs(hdir, exist_ok=True)
    verdicts = [verdict(r) for r in results]
    tally = {v: verdicts.count(v) for v in ("OK", "NG", "확인필요", "육안확인", "Mark-up 없음", "오류")}

    for r in results:
        rows = []
        for i in r["items"]:
            mk = f'<img loading="lazy" src="../{_rel(i["img_mk"], out_dir)}">' if i["img_mk"] else ""
            ab = f'<img loading="lazy" src="../{_rel(i["img_ab"], out_dir)}">' if i["img_ab"] else ""
            rows.append(
                f"<tr><td>{i['no']}</td><td>{html.escape(i['region'])}<br><small>{html.escape(i['field'])}</small></td>"
                f"<td>{html.escape(i['kind'])}<br><small>{html.escape(i['source'])}</small></td>"
                f"<td>{html.escape(i['old'])}</td><td><b>{html.escape(i['new'])}</b></td>"
                f"<td>{html.escape(i['asbuilt'])}</td><td>{_badge(i['status'])}<br><small>{html.escape(i['reason'])}</small></td>"
                f"<td><div class='imgs'><div><small>MARK-UP</small>{mk}</div><div><small>AS-BUILT</small>{ab}</div></div></td></tr>")
        err = f"<p class='s s-오류'>{html.escape(r['error'])}</p>" if r.get("error") else ""
        page = (f"<!doctype html><html lang='ko'><head><meta charset='utf-8'>"
                f"<meta name='viewport' content='width=device-width,initial-scale=1'><title>{html.escape(r['iso'])}</title>"
                f"<style>{CSS}</style></head><body><main><p><a href='../ISO_QC_Report.html'>← 목록</a></p>"
                f"<h1>{html.escape(r['iso'])} {_badge(verdict(r))}</h1>"
                f"<div class='sub'>MARK-UP: {html.escape(os.path.basename(r['markup']))} · AS-BUILT: "
                f"{html.escape(os.path.basename(r['asbuilt'] or ''))} · 정렬: {html.escape(r.get('align', ''))}</div>{err}"
                f"<div class='wrap'><table><tr><th>#</th><th>구분</th><th>유형</th><th>기존값</th><th>Mark-up 값</th>"
                f"<th>As-built 값</th><th>판정</th><th>비교 이미지 (초록=변경값, 빨강=기존값/불일치)</th></tr>"
                f"{''.join(rows) or '<tr><td colspan=8>Mark-up 없음</td></tr>'}</table></div></main></body></html>")
        with open(os.path.join(hdir, r["html"]), "w", encoding="utf-8") as f:
            f.write(page)

    trs = []
    for n, (r, v) in enumerate(zip(results, verdicts), 1):
        c = _counts(r["items"])
        trs.append(f"<tr><td>{n}</td><td><a href='html/{r['html']}'>{html.escape(r['iso'])}</a><br>"
                   f"<small>{html.escape(os.path.basename(r['markup']))}</small></td><td>{_badge(v)}</td>"
                   f"<td>{html.escape(_region_verdict(r['items'], 'Routing'))}</td>"
                   f"<td>{html.escape(_region_verdict(r['items'], 'Line Condition'))}</td>"
                   f"<td>{c[C.OK]}</td><td>{c[C.NG]}</td><td>{c[C.CHECK]}</td><td>{c[C.VISUAL]}</td>"
                   f"<td><small>{html.escape(r.get('error', ''))}</small></td></tr>")
    cards = "".join(f"<div class='card'>{html.escape(k)}<b>{v}</b></div>" for k, v in tally.items())
    idx = (f"<!doctype html><html lang='ko'><head><meta charset='utf-8'>"
           f"<meta name='viewport' content='width=device-width,initial-scale=1'><title>ISO QC Report</title>"
           f"<style>{CSS}</style></head><body><main><h1>CIMS Mark-up → As-built 반영 QC</h1>"
           f"<div class='sub'>{html.escape(meta.get('실행 시각', ''))} · 검토 {len(results)}건</div>"
           f"<div class='cards'>{cards}</div>"
           f"<input id='q' placeholder='ISO No / 판정 검색' oninput=\"for(const t of document.querySelectorAll('tbody tr'))"
           f"t.style.display=t.innerText.toUpperCase().includes(this.value.toUpperCase())?'':'none'\">"
           f"<div class='wrap'><table><thead><tr><th>No</th><th>ISO No</th><th>종합</th><th>1. Routing</th>"
           f"<th>2. Line Condition</th><th>반영</th><th>미반영</th><th>확인필요</th><th>육안확인</th><th>비고</th></tr>"
           f"</thead><tbody>{''.join(trs)}</tbody></table></div></main></body></html>")
    path = os.path.join(out_dir, "ISO_QC_Report.html")
    with open(path, "w", encoding="utf-8") as f:
        f.write(idx)
    return path
